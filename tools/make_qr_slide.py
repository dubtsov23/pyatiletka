# Слайд «Играем на телефоне» с QR-кодом — для вставки в скачанную презентацию.
# Запуск: python tools/make_qr_slide.py  (нужны python-pptx и Pillow)
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
URL = 'dubtsov23.github.io/pyatiletka'
NAVY, RED, PAPER, CARD, MUTED, LINE = (RGBColor.from_string(c) for c in ('14213D', 'A13A2C', 'F7F5F0', 'FFFDF8', '5F5A52', 'DDD6C8'))
SERIF, SANS = 'Georgia', 'Arial'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
s.background.fill.solid()
s.background.fill.fore_color.rgb = PAPER


def text(x, y, w, h, runs, size, color=NAVY, font=SANS, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    p = tf.paragraphs[0]
    p.alignment = align
    for r in (runs if isinstance(runs, list) else [(runs, color)]):
        run = p.add_run()
        run.text = r[0]
        run.font.size, run.font.name, run.font.bold = Pt(size), font, bold
        run.font.color.rgb = r[1]
    if spacing:
        p.line_spacing = spacing
    return tb


text(0.89, 0.62, 11.5, 0.35, 'ПРАКТИКА · ИГРА-СИМУЛЯТОР · 区域政策模拟游戏 · «ПЯТИЛЕТКА» · 五年规划', 12, RED, bold=True)
text(0.89, 0.98, 11.5, 0.8, 'Играем на телефоне · 用手机玩', 32, NAVY, SERIF, bold=True)

# QR на белой карточке
card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.89), Inches(1.95), Inches(4.6), Inches(5.0))
card.adjustments[0] = 0.05
card.fill.solid(); card.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
card.line.color.rgb = LINE
s.shapes.add_picture(os.path.join(ROOT, 'qr.png'), Inches(1.19), Inches(2.15), Inches(4.0), Inches(4.0))
text(0.89, 6.25, 4.6, 0.5, URL, 15, NAVY, bold=True, align=PP_ALIGN.CENTER)

steps = [
    ('Один телефон на команду из 3–4 человек', '每队一部手机，3–4 人'),
    ('Сканируйте QR-код камерой или в WeChat', '用相机或微信扫描二维码'),
    ('Язык RU / 中文 и название команды', '选择语言，输入队名'),
    ('Игра сама покажет событие и посчитает итоги', '游戏自动显示突发事件并计算结果'),
    ('Спикер называет Экономику, Разрыв, Индекс', '发言人报出经济总量、地区差距和指数'),
]
y = 2.0
for i, (ru, zh) in enumerate(steps, 1):
    num = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.0), Inches(y), Inches(0.52), Inches(0.52))
    num.fill.solid(); num.fill.fore_color.rgb = NAVY
    num.line.fill.background()
    tf = num.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(i); r.font.size = Pt(16); r.font.bold = True; r.font.name = SANS; r.font.color.rgb = PAPER
    tb = text(6.75, y - 0.04, 5.9, 0.8, ru, 18, NAVY, bold=True)
    p = tb.text_frame.add_paragraph()
    r = p.add_run(); r.text = zh; r.font.size = Pt(15); r.font.name = SANS; r.font.color.rgb = RED
    y += 0.92

bar = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.0), Inches(6.62), Inches(6.45), Inches(0.5))
bar.adjustments[0] = 0.2
bar.fill.solid(); bar.fill.fore_color.rgb = NAVY
bar.line.fill.background()
tf = bar.text_frame
tf.margin_left = Inches(0.2)
tf.vertical_anchor = MSO_ANCHOR.MIDDLE
p = tf.paragraphs[0]
r = p.add_run(); r.text = 'Нет сети — бумажный лист, правила те же · 没有网络就用纸质版'
r.font.size = Pt(13); r.font.name = SANS; r.font.color.rgb = PAPER

s.notes_slide.notes_text_frame.text = (
    'Игра «Пятилетка» на телефонах (вставить после слайда с правилами). Пока команды сканируют QR, '
    'назначьте роли: казначей, картограф, спикер. Игра сама показывает событие пятилетки и считает итоги, '
    'поэтому расчёт занимает меньше времени. В конце спикеры называют три числа; пять уровней регионов '
    'с экрана итогов вводятся в табло. Если сеть не работает — раздайте бумажные листы.')

out = os.path.join(ROOT, 'Пятилетка_слайд_QR.pptx')
prs.save(out)
print(out)
